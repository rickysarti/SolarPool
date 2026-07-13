import openpyxl
import sys

try:
    wb = openpyxl.load_workbook(r"E:\pagina cobertores precios.xlsx")
    
    print("=" * 100)
    print("EXCEL FILE: E:\pagina cobertores precios.xlsx")
    print("=" * 100)
    print(f"Sheet names: {wb.sheetnames}\n")
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n{'=' * 100}")
        print(f"SHEET: {sheet_name}")
        print(f"Dimensions: {ws.dimensions}")
        print(f"{'=' * 100}\n")
        
        for row_num in range(1, ws.max_row + 1):
            row_data = []
            for col_num in range(1, ws.max_column + 1):
                cell = ws.cell(row=row_num, column=col_num)
                if cell.value is not None:
                    row_data.append((cell.coordinate, cell.value))
            
            if row_data:
                print(f"Row {row_num}: {row_data}")

except Exception as e:
    print(f"Error: {e}", file=sys.stderr)
    import traceback
    traceback.print_exc()
